import os
import io
from datetime import datetime
from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse

# ── Library pembuatan dokumen ──────────────────────────────────────
# pip install python-docx openpyxl python-pptx Pillow
try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

try:
    from pptx import Presentation
    from pptx.util import Inches as PptInches, Pt as PptPt
    from pptx.dml.color import RGBColor as PptRGB
    from pptx.enum.text import PP_ALIGN
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    from PIL import Image as PILImage
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False

DATASET_DIR = "./static/datasets"

router = APIRouter(
    prefix="/api/documentation",
    tags=["Documentation & Report Generator"]
)

from ..schemas import ReportPayload

# Palet warna konsisten lintas format
C_PRIMARY = "1A56DB"   # biru
C_INK = "111827"       # hitam tulisan
C_MUTED = "6B7280"     # abu
C_BG_SOFT = "EFF6FF"
C_ROW_ALT = "F3F4F6"


def _list_images(folder_id: str):
    folder_path = os.path.join(DATASET_DIR, folder_id)
    if not os.path.exists(folder_path):
        return []
    return [f for f in sorted(os.listdir(folder_path))
            if f.lower().endswith(('.png', '.jpg', '.jpeg'))]


def _acq_mode(filename: str) -> str:
    """Tebak metode akuisisi dari pola nama berkas (lihat image-gathering)."""
    n = filename.lower()
    if n.startswith("tile_") or n.startswith(("img_2", "img_20")):
        return "Grid Scan (otomatis)"
    if "manual" in n:
        return "Manual"
    if n.startswith("img_input"):
        return "Input Images"
    if n.startswith(("yolo_", "edited_", "stitched_", "analysis_")):
        return "Hasil analisis"
    return "Tidak diketahui"


def _img_info(path: str):
    """(width, height, size_bytes). Resolusi via PIL; 0 kalau gagal."""
    size = os.path.getsize(path) if os.path.exists(path) else 0
    w = h = 0
    if PIL_AVAILABLE and os.path.exists(path):
        try:
            with PILImage.open(path) as im:
                w, h = im.size
        except Exception:
            pass
    return w, h, size


def _thumb_buf(path: str, max_wh):
    """Kembalikan BytesIO JPEG thumbnail, atau None."""
    if not (PIL_AVAILABLE and os.path.exists(path)):
        return None
    try:
        with PILImage.open(path) as im:
            im.thumbnail(max_wh)
            if im.mode in ("RGBA", "P", "LA"):
                im = im.convert("RGB")
            b = io.BytesIO()
            im.save(b, format="JPEG", quality=80)
            b.seek(0)
            return b
    except Exception:
        return None


def _limited(images, max_images):
    if max_images and max_images > 0:
        return images[:max_images]
    return images


def _folder_meta(payload: ReportPayload):
    class FolderMeta:
        id = payload.folder_id
        name = payload.folder_name or payload.folder_id
        object_type = payload.object_type or "Tidak diketahui"
        operator = payload.operator or (payload.author or "Tidak diketahui")
        date = payload.date or datetime.now().strftime("%Y-%m-%d")
    return FolderMeta()


# ─────────────────────────────────────────────────────────────────────
#  ENDPOINT
# ─────────────────────────────────────────────────────────────────────
@router.post("/generate")
async def generate_report(
    payload: ReportPayload,
    format: str = Query(..., description="Format dokumen: WORD, EXCEL, atau PPT"),
):
    folder = _folder_meta(payload)
    images = _list_images(payload.folder_id)
    folder_path = os.path.join(DATASET_DIR, payload.folder_id)
    if not os.path.exists(folder_path):
        raise HTTPException(status_code=404,
                            detail=f"Folder dataset '{payload.folder_id}' tidak ditemukan di server.")

    fmt = format.upper()
    if fmt == "WORD":
        return _generate_word(payload, folder, images, folder_path)
    if fmt == "EXCEL":
        return _generate_excel(payload, folder, images, folder_path)
    if fmt == "PPT":
        return _generate_ppt(payload, folder, images, folder_path)
    raise HTTPException(status_code=400, detail=f"Format tidak dikenali: {format}")


# ──────────────────────────────────────────────
#  Util docx
# ──────────────────────────────────────────────
def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _kv_table(doc, rows):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.style = "Table Grid"
    for k, v in rows:
        r = tbl.add_row().cells
        r[0].text = str(k)
        r[1].text = str(v)
        _shade(r[0], C_BG_SOFT)
        for p in r[0].paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(10)
        for p in r[1].paragraphs:
            for run in p.runs:
                run.font.size = Pt(10)
    tbl.columns[0].width = Inches(2.1)
    tbl.columns[1].width = Inches(4.2)
    return tbl


# ──────────────────────────────────────────────
#  GENERATOR 1: Word (.docx)
# ──────────────────────────────────────────────
def _generate_word(payload: ReportPayload, folder, images, folder_path):
    if not DOCX_AVAILABLE:
        raise HTTPException(status_code=500, detail="python-docx tidak terinstall di server.")
    try:
        title = payload.title or "Laporan Hasil Analisis Mikroskop"
        institution = payload.institution or "—"
        doc = Document()
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1.1)
            section.right_margin = Inches(1.1)

        base = doc.styles["Normal"]
        base.font.name = "Calibri"
        base.font.size = Pt(11)

        # ── Kop ──
        kop = doc.add_paragraph(institution.upper())
        kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
        kop.runs[0].font.size = Pt(11)
        kop.runs[0].font.bold = True
        kop.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        h = doc.add_heading(title, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.runs[0].font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

        sub = doc.add_paragraph(f"Objek penelitian: {folder.object_type}")
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.runs[0].font.size = Pt(11)
        sub.runs[0].font.italic = True
        if payload.doc_number:
            dn = doc.add_paragraph(f"No. Dokumen: {payload.doc_number}")
            dn.alignment = WD_ALIGN_PARAGRAPH.CENTER
            dn.runs[0].font.size = Pt(9)
            dn.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        rule = doc.add_paragraph("─" * 74)
        rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rule.runs[0].font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

        # ── Identitas laporan (tabel) ──
        doc.add_heading("A. Identitas Laporan", level=1)
        _kv_table(doc, [
            ("Nama Folder Dataset", folder.name),
            ("Jenis Objek", folder.object_type),
            ("Penyusun", payload.author or folder.operator),
            ("Operator Akuisisi", folder.operator),
            ("Institusi", payload.institution),
            ("Tanggal Pengambilan Data", folder.date),
            ("Jumlah Citra", str(len(images))),
            ("Waktu Pembuatan Dokumen", datetime.now().strftime("%d %B %Y, %H:%M:%S WIB")),
        ])
        doc.add_paragraph()

        # ── Abstrak ──
        doc.add_heading("B. Abstrak", level=1)
        abstract = payload.abstract.strip() or (
            f'Laporan ini mendokumentasikan {len(images)} citra mikroskopik objek "{folder.object_type}" '
            f'dari folder dataset "{folder.name}". Data diakuisisi menggunakan sistem Digital Microscopy '
            f'BRIN/UNDIP berbasis Jetson Orin Nano dengan kamera IMX477 dan meja CNC terkontrol GRBL, '
            f'lalu diproses dengan pipeline Computer Vision (OpenCV) serta inferensi YOLO untuk '
            f'segmentasi dan perhitungan koloni.'
        )
        doc.add_paragraph(abstract)
        doc.add_paragraph()

        # ── Metodologi ──
        if payload.include_technical:
            doc.add_heading("C. Metodologi Akuisisi & Pemrosesan", level=1)
            modes = {}
            widths = []
            for name in images:
                modes[_acq_mode(name)] = modes.get(_acq_mode(name), 0) + 1
                w, hgt, _ = _img_info(os.path.join(folder_path, name))
                if w and hgt:
                    widths.append((w, hgt))
            res_txt = "tidak terbaca"
            if widths:
                wmin = min(w for w, _ in widths)
                wmax = max(w for w, _ in widths)
                hmin = min(h for _, h in widths)
                hmax = max(h for _, h in widths)
                res_txt = (f"{wmin}×{hmin} px"
                           if (wmin, hmin) == (wmax, hmax)
                           else f"{wmin}×{hmin} – {wmax}×{hmax} px")
            _kv_table(doc, [
                ("Perangkat Akuisisi", "Jetson Orin Nano + Sony IMX477 (nvarguscamerasrc)"),
                ("Kendali Gerak", "CNC 3-sumbu, firmware GRBL (G-code)"),
                ("Rentang Resolusi Citra", res_txt),
                ("Distribusi Metode", ", ".join(f"{k}: {v}" for k, v in modes.items()) or "-"),
                ("Pipeline Pemrosesan", "OpenCV (klasik) + YOLO Segmentation TensorRT INT8 di Edge"),
            ])
            doc.add_paragraph()

        # ── Rincian berkas ──
        if payload.include_table and images:
            doc.add_heading("D. Rincian Berkas Citra", level=1)
            tbl = doc.add_table(rows=1, cols=5)
            tbl.style = "Table Grid"
            for i, htext in enumerate(["No", "Nama Berkas", "Metode", "Ukuran", "Resolusi"]):
                c = tbl.rows[0].cells[i]
                c.text = htext
                _shade(c, C_PRIMARY)
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(9)
            for idx, name in enumerate(images, start=1):
                w, hgt, size = _img_info(os.path.join(folder_path, name))
                cells = tbl.add_row().cells
                vals = [str(idx), name, _acq_mode(name),
                        f"{size/1024:.0f} KB" if size else "-",
                        f"{w}×{hgt}" if w else "-"]
                for i, val in enumerate(vals):
                    cells[i].text = val
                    for p in cells[i].paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(8.5)
                    if idx % 2 == 0:
                        _shade(cells[i], C_ROW_ALT)
            doc.add_paragraph()

        # ── Galeri ──
        if payload.include_gallery and images:
            shown = _limited(images, payload.max_images)
            doc.add_heading("E. Galeri Citra Mikroskopik", level=1)
            doc.add_paragraph(
                f"Menampilkan {len(shown)} dari {len(images)} citra dalam folder ini."
                + ("" if len(shown) == len(images) else " (dibatasi oleh pengaturan laporan)")
            )
            for i, img_name in enumerate(shown):
                buf = _thumb_buf(os.path.join(folder_path, img_name), (820, 620))
                try:
                    if buf is not None:
                        doc.add_picture(buf, width=Inches(4.6))
                    else:
                        doc.add_picture(os.path.join(folder_path, img_name), width=Inches(4.6))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    continue
                cap = doc.add_paragraph(f"Gambar {i+1}. {img_name}")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].font.size = Pt(9)
                cap.runs[0].font.italic = True

        # ── Kesimpulan ──
        doc.add_heading("F. Kesimpulan", level=1)
        conclusion = payload.conclusion.strip() or (
            f"Sistem berhasil mendokumentasikan {len(images)} citra mikroskopik dari folder "
            f'"{folder.name}". Dataset ini siap digunakan sebagai lampiran resmi laporan penelitian, '
            f"publikasi, atau tugas akhir, serta sebagai bahan pelatihan/evaluasi model deteksi objek."
        )
        doc.add_paragraph(conclusion)
        doc.add_paragraph()
        doc.add_paragraph()

        # ── Blok tanda tangan ──
        place_date = f"{'Semarang'}, {datetime.now().strftime('%d %B %Y')}"
        pd = doc.add_paragraph(place_date)
        pd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        sig = doc.add_table(rows=1, cols=2)
        sig.alignment = WD_TABLE_ALIGNMENT.CENTER
        left, right = sig.rows[0].cells
        left.text = "Disusun oleh,\n\n\n\n" + (payload.author or folder.operator)
        right.text = "Mengetahui,\n\n\n\n" + (payload.supervisor or "________________________")
        for cell in (left, right):
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10)

        footer = doc.add_paragraph(
            f"Dokumen ini dikompilasi otomatis oleh Digital Microscopy System — "
            f"{payload.institution} — {datetime.now().year}"
        )
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.runs[0].font.size = Pt(8)
        footer.runs[0].font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

        out = io.BytesIO()
        doc.save(out)
        out.seek(0)
        safe = (title or "laporan").replace(" ", "_")
        return StreamingResponse(
            out,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{safe}.docx"'},
        )
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        print(f"ERROR IN _generate_word: {e}", flush=True)
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Gagal menyusun Word: {e}")


# ──────────────────────────────────────────────
#  GENERATOR 2: Excel (.xlsx)
# ──────────────────────────────────────────────
def _generate_excel(payload: ReportPayload, folder, images, folder_path):
    if not XLSX_AVAILABLE:
        raise HTTPException(status_code=500, detail="openpyxl tidak terinstall di server.")

    title = payload.title or "Laporan Hasil Analisis Mikroskop"
    wb = openpyxl.Workbook()

    hdr_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── Sheet 1: Laporan ──
    ws = wb.active
    ws.title = "Laporan"
    ws.merge_cells("A1:F1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", bold=True, size=15, color=C_PRIMARY)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].fill = PatternFill("solid", fgColor=C_BG_SOFT)
    ws.row_dimensions[1].height = 32
    ws.merge_cells("A2:F2")
    ws["A2"] = payload.institution
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].font = Font(italic=True, color=C_MUTED)

    meta = [
        ("Nama Folder", folder.name),
        ("Jenis Objek", folder.object_type),
        ("Penyusun", payload.author or folder.operator),
        ("Operator Akuisisi", folder.operator),
        ("Tanggal Data", folder.date),
        ("No. Dokumen", payload.doc_number or "-"),
        ("Total Gambar", len(images)),
        ("Waktu Laporan", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    for i, (k, v) in enumerate(meta, start=4):
        ws.cell(i, 1, k).font = Font(bold=True, size=10)
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=C_ROW_ALT)
        ws.cell(i, 2, str(v))

    row = 4 + len(meta) + 1
    if payload.abstract.strip():
        ws.cell(row, 1, "Abstrak").font = Font(bold=True, size=10)
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 2, end_column=6)
        ac = ws.cell(row, 2, payload.abstract.strip())
        ac.alignment = Alignment(wrap_text=True, vertical="top")
        row += 4
    if payload.conclusion.strip():
        ws.cell(row, 1, "Kesimpulan").font = Font(bold=True, size=10)
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 2, end_column=6)
        cc = ws.cell(row, 2, payload.conclusion.strip())
        cc.alignment = Alignment(wrap_text=True, vertical="top")
        row += 4

    # Tabel berkas
    hr = row + 1
    headers = ["No", "Nama Berkas", "Metode Akuisisi", "Ukuran (KB)", "Resolusi (px)", "Ekstensi"]
    for ci, htext in enumerate(headers, start=1):
        c = ws.cell(hr, ci, htext)
        c.font = hdr_font
        c.fill = PatternFill("solid", fgColor=C_PRIMARY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    ws.row_dimensions[hr].height = 22

    total_bytes = 0
    for i, name in enumerate(images):
        dr = hr + 1 + i
        w, hgt, size = _img_info(os.path.join(folder_path, name))
        total_bytes += size
        ext = os.path.splitext(name)[1].upper().replace(".", "")
        vals = [i + 1, name, _acq_mode(name),
                round(size / 1024) if size else 0,
                f"{w}×{hgt}" if w else "-", ext]
        fill = PatternFill("solid", fgColor="FFFFFF" if i % 2 == 0 else C_ROW_ALT)
        for ci, val in enumerate(vals, start=1):
            c = ws.cell(dr, ci, val)
            c.border = border
            c.fill = fill
            c.alignment = Alignment(vertical="center")

    last = hr + len(images)
    ws.freeze_panes = ws.cell(hr + 1, 1)
    if images:
        ws.auto_filter.ref = f"A{hr}:F{last}"
    for ci, wdt in enumerate([5, 42, 22, 13, 16, 11], start=1):
        ws.column_dimensions[get_column_letter(ci)].width = wdt

    # ── Sheet 2: Ringkasan ──
    ws2 = wb.create_sheet("Ringkasan")
    ws2["A1"] = "Ringkasan Dataset"
    ws2["A1"].font = Font(bold=True, size=13, color=C_PRIMARY)
    modes = {}
    for n in images:
        modes[_acq_mode(n)] = modes.get(_acq_mode(n), 0) + 1
    srows = [
        ["Parameter", "Nilai"],
        ["Total Gambar", len(images)],
        ["Total Ukuran", f"{total_bytes/1_048_576:.2f} MB"],
        ["Objek Penelitian", folder.object_type],
        ["Penyusun", payload.author or folder.operator],
        ["Operator", folder.operator],
        ["Tanggal", folder.date],
        ["Sistem Akuisisi", "Jetson Orin Nano + IMX477 + CNC GRBL"],
    ]
    for k, v in modes.items():
        srows.append([f"  • {k}", v])
    for ri, r in enumerate(srows, start=3):
        for ci, val in enumerate(r, start=1):
            c = ws2.cell(ri, ci, val)
            c.border = border
            if ri == 3:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=C_PRIMARY)
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 34

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    safe = (title or "laporan").replace(" ", "_")
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe}.xlsx"'},
    )


# ──────────────────────────────────────────────
#  GENERATOR 3: PowerPoint (.pptx)
# ──────────────────────────────────────────────
def _generate_ppt(payload: ReportPayload, folder, images, folder_path):
    if not PPTX_AVAILABLE:
        raise HTTPException(status_code=500, detail="python-pptx tidak terinstall di server.")

    title = payload.title or "Laporan Hasil Analisis Mikroskop"
    institution = payload.institution or "—"
    prs = Presentation()
    prs.slide_width = PptInches(13.33)
    prs.slide_height = PptInches(7.5)
    blank = prs.slide_layouts[6]

    def rgb(h):
        return PptRGB(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    BLUE = rgb(C_PRIMARY)
    WHITE = rgb("FFFFFF")
    DARK = rgb(C_INK)
    GRAY_T = rgb(C_MUTED)
    LIGHT = rgb("93C5FD")

    def bg(slide, hexc):
        # dipanggil PALING AWAL di tiap slide -> otomatis jadi lapisan paling belakang
        sh = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(hexc)
        sh.line.fill.background()

    def rect(slide, x, y, w, h, hexc):
        sh = slide.shapes.add_shape(1, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(hexc)
        sh.line.fill.background()

    def tb(slide, x, y, w, h, text, size, bold=False, color=None, align="center"):
        txb = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
        tf = txb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        run = p.runs[0]
        run.font.size = PptPt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                       "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.CENTER)

    # Slide 1: Cover
    s1 = prs.slides.add_slide(blank)
    bg(s1, "0F172A")
    tb(s1, 1, 0.9, 11.33, 0.5, institution.upper(), 12, bold=True, color=LIGHT)
    rect(s1, 1, 3.35, 6.5, 0.06, C_PRIMARY)
    tb(s1, 1, 1.7, 11.33, 1.6, title, 34, bold=True, color=WHITE, align="left")
    tb(s1, 1, 3.6, 11.33, 1.4,
       f"Objek: {folder.object_type}\nFolder: {folder.name}\nPenyusun: {payload.author or folder.operator}"
       f"   |   Tanggal data: {folder.date}", 13, color=LIGHT, align="left")
    if payload.doc_number:
        tb(s1, 1, 6.2, 11.33, 0.4, f"No. Dokumen: {payload.doc_number}", 10, color=GRAY_T, align="left")
    tb(s1, 1, 6.8, 11.33, 0.4,
       f"Digital Microscopy System — {datetime.now().strftime('%d %B %Y')}", 9, color=GRAY_T, align="left")

    # Slide 2: Identitas
    s2 = prs.slides.add_slide(blank)
    bg(s2, "F9FAFB")
    rect(s2, 0, 0, 13.33, 1.05, C_PRIMARY)
    tb(s2, 0.4, 0.18, 12, 0.7, "Identitas Laporan", 22, bold=True, color=WHITE, align="left")
    rows = [
        ("Nama Folder", folder.name),
        ("Jenis Objek", folder.object_type),
        ("Penyusun", payload.author or folder.operator),
        ("Operator Akuisisi", folder.operator),
        ("Institusi", payload.institution),
        ("Tanggal Data", folder.date),
        ("Total Gambar", str(len(images))),
        ("Perangkat", "Jetson Orin Nano + IMX477 + CNC GRBL"),
    ]
    for i, (k, v) in enumerate(rows):
        y = 1.35 + i * 0.72
        tb(s2, 0.6, y, 4.4, 0.55, k, 12, bold=True, color=DARK, align="left")
        tb(s2, 5.2, y, 7.6, 0.55, str(v), 12, color=BLUE, align="left")

    # Slide 3: Abstrak & Kesimpulan
    s3 = prs.slides.add_slide(blank)
    bg(s3, "FFFFFF")
    rect(s3, 0, 0, 13.33, 1.05, C_PRIMARY)
    tb(s3, 0.4, 0.18, 12, 0.7, "Abstrak & Kesimpulan", 22, bold=True, color=WHITE, align="left")
    abstract = payload.abstract.strip() or (
        f"{len(images)} citra mikroskopik objek \"{folder.object_type}\" dari folder \"{folder.name}\", "
        f"diakuisisi dengan sistem Digital Microscopy BRIN/UNDIP (Jetson Orin Nano, IMX477, CNC GRBL) "
        f"dan diproses dengan pipeline OpenCV + YOLO Segmentation di Edge.")
    conclusion = payload.conclusion.strip() or (
        f"Dataset \"{folder.name}\" terdokumentasi lengkap dan siap dipakai sebagai lampiran "
        f"laporan penelitian / tugas akhir maupun bahan pelatihan model deteksi.")
    tb(s3, 0.7, 1.4, 12, 0.4, "Abstrak", 13, bold=True, color=DARK, align="left")
    tb(s3, 0.7, 1.9, 12, 2.2, abstract, 12, color=DARK, align="left")
    tb(s3, 0.7, 4.3, 12, 0.4, "Kesimpulan", 13, bold=True, color=DARK, align="left")
    tb(s3, 0.7, 4.8, 12, 2.2, conclusion, 12, color=DARK, align="left")

    # Slide galeri
    if payload.include_gallery:
        shown = _limited(images, payload.max_images)
        for cs in range(0, len(shown), 2):
            chunk = shown[cs:cs + 2]
            si = prs.slides.add_slide(blank)
            bg(si, "111827")
            tb(si, 0.3, 0.15, 12, 0.5,
               f"Galeri Citra ({cs+1}–{min(cs+2, len(shown))} / {len(images)})",
               14, bold=True, color=LIGHT, align="left")
            for j, img_name in enumerate(chunk):
                xp = 0.3 + j * 6.53
                buf = _thumb_buf(os.path.join(folder_path, img_name), (900, 700))
                try:
                    if buf is not None:
                        si.shapes.add_picture(buf, PptInches(xp), PptInches(0.8), width=PptInches(6.2))
                    else:
                        si.shapes.add_picture(os.path.join(folder_path, img_name),
                                              PptInches(xp), PptInches(0.8), width=PptInches(6.2))
                    tb(si, xp, 6.75, 6.2, 0.35, img_name, 8, color=GRAY_T, align="left")
                except Exception:
                    pass

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    safe = (title or "laporan").replace(" ", "_")
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{safe}.pptx"'},
    )
